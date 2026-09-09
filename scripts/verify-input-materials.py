"""Offline regression: user headers, asset snapshots, refresh and flow documents."""
import json
import os
import sys
import tempfile
import base64
import runpy
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from pangea_agent.assets import import_asset, prepare_asset_extraction, freeze_asset_inputs, load_asset, review_asset
from pangea_agent.documents.coverage_input import read_input, gap_records, freeze_input, prepare_coverage


def verify():
    with tempfile.TemporaryDirectory(prefix="pangea-input-regression-") as directory:
        root = Path(directory)
        report = root / "函数覆盖率.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(['特性', '模块', '代码 路径', '函数名', '是否覆盖', None, '覆盖次数'])
        sheet.append(['synthetic', 'tls', r'src\tls.c', 'handshake', '否', None, 0])
        sheet.append(['synthetic', 'tls', 'src/tls.c', 'covered', '是', None, 3])
        sheet.append(['synthetic', 'tls', 'src/tls.c', 'unknown', '', None, '-'])
        sheet.append(['synthetic', 'tls', 'src/tls.c', 'conflict', '是', None, 0])
        sheet.append(['synthetic', 'tls', None, 'unlocated', '否', None, 0])
        other = book.create_sheet('another module')
        other.append(['source', 'file_path', 'kind', 'count', 'line', 'block', 'branch'])
        other.append(['summary', 'src/io.c', 'branch', 0, 10, 0, 1])
        book.create_sheet('说明').append(['not a coverage table'])
        book.save(report)
        data = read_input(report)
        assert len(data['records']) == 6
        assert len(data['unknown_records']) == 2
        assert len(gap_records(data)) == 3
        assert gap_records(data)[0]['file_path'] == r'src\tls.c'
        assert gap_records(data)[-1]['location_status'] == 'unresolved'
        asset = import_asset(str(root), str(report), 'coverage')
        run = root / 'run'
        run.mkdir()
        manifest = freeze_asset_inputs(str(root), run, 'run', [asset.asset_id])
        freeze_input({'kind': 'asset', 'asset_id': asset.asset_id}, run, manifest)
        frozen = (run / 'inputs/coverage/combined.json').read_bytes()
        parsed = read_input(run / 'inputs/coverage/combined.json')
        assert [(r['file_path'], r['kind'], r['raw']) for r in gap_records(parsed)] == [(r['file_path'], r['kind'], r['raw']) for r in gap_records(read_input(Path(asset.result_path)))]
        previous = asset.result_path
        prepare_asset_extraction(str(root), asset.asset_id)
        assert Path(previous).is_file() and load_asset(str(root), asset.asset_id).result_path != previous
        assert (run / 'inputs/coverage/combined.json').read_bytes() == frozen
        from pangea_agent.skill_runs import create_skill_run
        repository = root / 'repositories/example'
        repository.mkdir(parents=True)
        (repository / 'sample.c').write_text('int sample(void) { return 0; }', encoding='utf-8')
        request_path = root / 'asset-request.json'
        request_path.write_text(json.dumps({'request_version':'2.0', 'data_root':str(root), 'repository':'example',
            'target':'synthetic', 'source_scope':[], 'asset_ids':[], 'scenario':'coverage-analysis', 'mode':'speed',
            'coverage_input':{'kind':'asset','asset_id':asset.asset_id}}), encoding='utf-8')
        created = create_skill_run(str(request_path))
        assert prepare_coverage(Path(created['run_root']))['total'] == 3
        with patch('pangea_agent.assets.read_input', side_effect=ValueError('synthetic parse failure')):
            try:
                prepare_asset_extraction(str(root), asset.asset_id)
                raise AssertionError('expected parse error')
            except ValueError:
                pass
        assert load_asset(str(root), asset.asset_id).status == 'available'

        document = root / 'history.md'
        document.write_text('# DEFECT-1\nSynthetic observed fault', encoding='utf-8')
        history = import_asset(str(root), str(document), 'historical_defect')
        review_asset(str(root), history.asset_id, 'approve')
        prepare_asset_extraction(str(root), history.asset_id)
        assert load_asset(str(root), history.asset_id).review_status == 'approved'
        with patch('pangea_agent.assets.extract_document', side_effect=ValueError('synthetic document failure')):
            try:
                prepare_asset_extraction(str(root), history.asset_id)
            except ValueError:
                pass
        assert load_asset(str(root), history.asset_id).status == 'available'

        from docx import Document
        image = root / 'pixel.png'
        image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aN1sAAAAASUVORK5CYII='))
        design = Document()
        design.add_paragraph('REQ-1 before table')
        design.add_table(rows=1, cols=1).cell(0, 0).text = 'design constraint'
        design.add_paragraph('REQ-2 after table')
        design.add_picture(str(image))
        design_path = root / 'design.docx'
        design.save(design_path)
        design_asset = import_asset(str(root), str(design_path), 'design')
        normalized = Path(design_asset.normalized_text_path).read_text(encoding='utf-8')
        assert normalized.index('REQ-1') < normalized.index('design constraint') < normalized.index('REQ-2')
        doc_run = root / 'document-run'
        manifest = freeze_asset_inputs(str(root), doc_run, 'document-run', [design_asset.asset_id])
        attachments = manifest['assets'][0]['attachments']
        assert len(attachments) == 1 and Path(attachments[0]['attachment_path']).is_file()
        assert Path(attachments[0]['attachment_path']).is_relative_to(doc_run)

        from pangea_agent.skills import SOURCE_ROOT
        for package in ['codetalks-skill', 'codetalks-coverage-skill']:
            enrich = runpy.run_path(str(SOURCE_ROOT.parent / package / 'scripts/flow_documents.py'))['enrich_flows']
            flow_path = doc_run / '活文档/流程讲解/流程-FLOW-1.md'
            flow_path.parent.mkdir(parents=True, exist_ok=True)
            flow_path.write_text('```pangea-flow\n' + json.dumps({'flow_id':'FLOW-1', 'mainline_steps':[{'step_id':'S1','title':'Open'}], 'branches':[]}) + '\n```', encoding='utf-8')
            flow = {'flow_id':'FLOW-1', 'title':'Synthetic', 'document_path':flow_path.relative_to(doc_run).as_posix()}
            assert enrich({'business_flows':[flow]}, doc_run)['business_flows'][0]['mainline_steps'][0]['step_id'] == 'S1'
            flow_path.write_text('Partial write: ```pangea-flow', encoding='utf-8')
            assert enrich({'business_flows':[flow]}, doc_run)['flow_document_warnings']

        query = root / 'query-run'
        folder = query / 'inputs/coverage'
        folder.mkdir(parents=True)
        skill = query / 'inputs/coverage-query-skill/scripts'
        skill.mkdir(parents=True)
        (skill / 'coverage_query.py').write_text('''import json,sys
args=sys.argv
product=args[args.index('--product')+1]
version=args[args.index('--version')+1]
print(json.dumps({'status':'no_data','sources':[], 'warnings':[], 'query_resolution':{'product':product,'version':version}}))
''', encoding='utf-8')
        original = {'kind':'query','query':{'product':'Synthetic Product','c_version':'R1','module':'tls','b_version':''}, 'skill_root':str(skill.parent)}
        (folder / 'input.json').write_text(json.dumps(original), encoding='utf-8')
        prepare_coverage(query)
        old = (folder / 'combined.json').read_bytes()
        updated = prepare_coverage(query, refresh={**original['query'], 'c_version':'Synthetic Product R2'})
        assert updated['query_resolution']['version'] == 'Synthetic Product R2'
        assert updated['query_resolution']['product'] == 'Synthetic Product'
        archives = list((folder / 'acquisitions').glob('*/combined.json'))
        assert len(archives) == 1 and archives[0].read_bytes() == old
        assert not (folder / 'acquisition.lock').exists()
        print('PASS: headers, unknown/conflicting counts, locations, asset reuse, preserved snapshots/reviews, query refresh')


if __name__ == '__main__':
    verify()
